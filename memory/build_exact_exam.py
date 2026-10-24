import docx
from docx.shared import Cm, Inches, Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

kop_path = r'C:\Users\yonat\Desktop\Soal Ujian Semester I T.P.2026-2027\SD SWASTA METHODIST.docx'
backup_path = r'C:\Users\yonat\.claude\projects\C--Users-yonat\memory\backups\Agama 1 Rio OK.raw_backup.docx'
target_path = r'C:\Users\yonat\Desktop\Soal Ujian Semester I T.P.2026-2027\Agama\Agama 1 Rio OK.docx'

doc_kop = docx.Document(kop_path)
tbl_xml = doc_kop.tables[0]._tbl.xml

doc = docx.Document(backup_path)

# Helper to create continuous sectPr
def make_sectPr(cols=1):
    col_str = f'<w:cols {nsdecls("w")} w:num="{cols}" w:space="708"/>' if cols > 1 else f'<w:cols {nsdecls("w")} w:space="708"/>'
    return parse_xml(f'''
    <w:sectPr {nsdecls("w")}>
      <w:type w:val="continuous"/>
      <w:pgSz w:w="12240" w:h="18720" w:code="14"/>
      <w:pgMar w:top="283" w:right="567" w:bottom="283" w:left="567" w:header="709" w:footer="709" w:gutter="0"/>
      {col_str}
      <w:docGrid w:linePitch="360"/>
    </w:sectPr>
    ''')

# Helper to create empty paragraph
def make_empty_p():
    return parse_xml(f'''
    <w:p {nsdecls("w")}>
      <w:pPr>
        <w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>
      </w:pPr>
    </w:p>
    ''')

# 1. Remove old header lines
p0 = doc.paragraphs[0]._p
p1 = doc.paragraphs[1]._p
p0.getparent().remove(p0)
p1.getparent().remove(p1)

# Clean tabs and trailing spaces in text
for p in doc.paragraphs:
    t = p.text
    if '\t' in t or '  ' in t:
        cleaned = t.replace('\t', '').strip()
        if cleaned:
            p.text = cleaned

# 2. Insert Kop Table at the top
new_tbl = parse_xml(tbl_xml)
doc.paragraphs[0]._p.addprevious(new_tbl)

# Fill Kop Table
t = doc.tables[0]
for r_i, val in [(0, 'Agama'), (1, 'Selasa, 1 Desember 2026'), (2, 'IA SD')]:
    cell_p = t.cell(r_i, 4).paragraphs[0]
    cell_p.text = ''
    run = cell_p.add_run(val)
    run.font.name = 'Tahoma'
    run.font.size = Pt(7)
    run.font.bold = True

# 3. Section 0: Kop Surat (1 col, continuous)
# Insert separator paragraph holding Section 0 sectPr
sep_p_xml = f'''
<w:p {nsdecls("w")}>
  <w:pPr>
    <w:spacing w:before="0" w:after="0" w:line="20" w:lineRule="exact"/>
    <w:rPr>
      <w:sz w:val="2"/>
      <w:szCs w:val="2"/>
    </w:rPr>
    <w:sectPr>
      <w:type w:val="continuous"/>
      <w:pgSz w:w="12240" w:h="18720" w:code="14"/>
      <w:pgMar w:top="283" w:right="567" w:bottom="283" w:left="567" w:header="709" w:footer="709" w:gutter="0"/>
      <w:cols w:num="1" w:space="708"/>
      <w:docGrid w:linePitch="360"/>
    </w:sectPr>
  </w:pPr>
</w:p>
'''
doc.paragraphs[0]._p.addprevious(parse_xml(sep_p_xml))

# doc.paragraphs[0] is sep_p (1-pt)
# doc.paragraphs[1] is 'Pilihan Ganda CP 1 (100%)' -> rapat dengan kop
doc.paragraphs[1].paragraph_format.space_before = Pt(0)
doc.paragraphs[1].paragraph_format.space_after = Pt(2)

# Ensure enter before 'CP 2 (100%)' and 'CP 3 (100%)'
for idx in range(len(doc.paragraphs)):
    p = doc.paragraphs[idx]
    if p.text.strip() in ['CP 2 (100%)', 'CP 3 (100%)']:
        prev_p = doc.paragraphs[idx - 1]
        if prev_p.text.strip() != '':
            p._p.addprevious(make_empty_p())

# Find the end of Pilihan Ganda (last option before Isian)
# and find Isian heading
isian_p = None
for p in doc.paragraphs:
    if 'Isian CP 4' in p.text:
        isian_p = p
        break

# The paragraph right before isian_p (or before empty line before isian_p) ends Section 1 (2 cols)
# Let's inspect paragraphs preceding isian_p
curr = isian_p._p.getprevious()
# If previous is empty paragraph, find the real content before it
while curr is not None and curr.text is not None and curr.text.strip() == '':
    curr = curr.getprevious()

# curr is now the last option of PG (e.g. 'Yakub, Esau, dan Yusuf')
# It will hold sectPr(cols=2) to end the 2-col Pilihan Ganda section!
curr.get_or_add_pPr().append(make_sectPr(cols=2))

# Now between PG and Isian:
# "(ada enter sekali 1 kolom ) lalu lanjut isian 2 kolom"
# We insert an empty paragraph in 1 kolom!
enter_1col = make_empty_p()
# The Isian heading is in 1 kolom:
# isian_p will hold sectPr(cols=1) so that the enter + Isian heading are in 1 col,
# and right after isian_p, the next section starts as 2 cols!
isian_p._p.addprevious(enter_1col)
isian_p.paragraph_format.space_before = Pt(0)
isian_p.paragraph_format.space_after = Pt(2)

# End the 1-col Isian heading section with sectPr(cols=1)
isian_p._p.get_or_add_pPr().append(make_sectPr(cols=1))

# Now Isian questions (5 questions): they will be in 2 cols!
# Find Esai heading
esai_p = None
for p in doc.paragraphs:
    if 'Esai (60%)' in p.text:
        esai_p = p
        break

# The last isian question is right before esai_p
curr_isian = esai_p._p.getprevious()
while curr_isian is not None and curr_isian.text is not None and curr_isian.text.strip() == '':
    curr_isian = curr_isian.getprevious()

# curr_isian holds sectPr(cols=2) to end the 2-col Isian section!
curr_isian.get_or_add_pPr().append(make_sectPr(cols=2))

# Now Esai: "baru essay 1 kolom"
# Insert enter before Esai
enter_esai = make_empty_p()
esai_p._p.addprevious(enter_esai)
esai_p.paragraph_format.space_before = Pt(0)
esai_p.paragraph_format.space_after = Pt(2)

# Document-level sectPr (ends the document): 1 col, continuous!
body_sectPr = doc.sections[-1]._sectPr
for c in body_sectPr.xpath('.//w:cols'):
    c.getparent().remove(c)
body_sectPr.append(parse_xml(f'<w:cols {nsdecls("w")} w:num="1" w:space="708"/>'))
if not body_sectPr.xpath('.//w:type'):
    body_sectPr.insert(0, parse_xml(f'<w:type {nsdecls("w")} w:val="continuous"/>'))
else:
    body_sectPr.xpath('.//w:type')[0].set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', 'continuous')

# Ensure page size and margins
doc.sections[-1].page_width = Inches(8.5)
doc.sections[-1].page_height = Inches(13.0)
doc.sections[-1].top_margin = Cm(0.5)
doc.sections[-1].bottom_margin = Cm(0.5)
doc.sections[-1].left_margin = Cm(1.0)
doc.sections[-1].right_margin = Cm(1.0)

doc.save(target_path)
print("SUCCESS: Full custom layout implemented!")
