import docx
from docx.shared import Cm, Inches, Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

kop_path = r'C:\Users\yonat\Desktop\Soal Ujian Semester I T.P.2026-2027\SD SWASTA METHODIST.docx'
target_path = r'C:\Users\yonat\Desktop\Soal Ujian Semester I T.P.2026-2027\Agama\Agama 1 Rio OK.docx'

doc_kop = docx.Document(kop_path)
tbl_xml = doc_kop.tables[0]._tbl.xml

doc = docx.Document(target_path)

# If table already exists, remove it first for clean rebuild
while doc.tables:
    t_elem = doc.tables[0]._tbl
    t_elem.getparent().remove(t_elem)

# Insert Kop table before the very first paragraph
new_tbl = parse_xml(tbl_xml)
doc.paragraphs[0]._p.addprevious(new_tbl)

# Update kop table cells
t = doc.tables[0]
for r_i, val in [(0, 'Agama'), (1, 'Selasa, 1 Desember 2026'), (2, 'IA SD')]:
    p = t.cell(r_i, 4).paragraphs[0]
    p.text = ''
    run = p.add_run(val)
    run.font.name = 'Tahoma'
    run.font.size = Pt(7)
    run.font.bold = True

# Paragraph 0 after the table should hold the Section 0 section break (1-column kop)
# We insert a separator paragraph right after the table
sep_xml = f'''
<w:p {nsdecls("w")}>
  <w:pPr>
    <w:spacing w:before="0" w:after="0" w:line="20" w:lineRule="exact"/>
    <w:rPr>
      <w:sz w:val="2"/>
      <w:szCs w:val="2"/>
    </w:rPr>
    <w:sectPr>
      <w:pgSz w:w="12240" w:h="18720" w:code="14"/>
      <w:pgMar w:top="283" w:right="567" w:bottom="283" w:left="567" w:header="709" w:footer="709" w:gutter="0"/>
      <w:cols w:num="1" w:space="708"/>
      <w:docGrid w:linePitch="360"/>
    </w:sectPr>
  </w:pPr>
</w:p>
'''
doc.paragraphs[0]._p.addprevious(parse_xml(sep_xml))

# Now doc.sections:
# Section 0 is 1 column (Kop Surat).
# Section 1 is 2 columns (Soal), Continuous.
s1 = doc.sections[1]
s1.start_type = docx.enum.section.WD_SECTION_START.CONTINUOUS
s1.page_width = Inches(8.5)
s1.page_height = Inches(13.0)
s1.top_margin = Cm(0.5)
s1.bottom_margin = Cm(0.5)
s1.left_margin = Cm(1.0)
s1.right_margin = Cm(1.0)

cols_elem = s1._sectPr.xpath('.//w:cols')
if cols_elem:
    cols_elem[0].set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}num', '2')
    cols_elem[0].set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}space', '708')
else:
    s1._sectPr.append(parse_xml(f'<w:cols {nsdecls("w")} w:num="2" w:space="708"/>'))

# First title after kop: must be tight with kop
# doc.paragraphs[0] is the separator paragraph holding Section 0 sectPr
# doc.paragraphs[1] is the first title
first_title = doc.paragraphs[1]
first_title.paragraph_format.space_before = Pt(0)
first_title.paragraph_format.space_after = Pt(2)

# Headings that must have an enter above them:
# CP 2 (100%), CP 3 (100%), Isian CP 4 (40%), Esai (60%)
target_headings = ['CP 2 (100%)', 'CP 3 (100%)', 'Isian CP 4 (40%)', 'Esai (60%)']

for idx, p in enumerate(list(doc.paragraphs)):
    t_text = p.text.strip()
    if any(h == t_text for h in target_headings):
        prev_p = doc.paragraphs[idx - 1] if idx > 0 else None
        if prev_p is not None and prev_p.text.strip() != '':
            blank_xml = f'''
            <w:p {nsdecls("w")}>
              <w:pPr>
                <w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>
              </w:pPr>
            </w:p>
            '''
            p._p.addprevious(parse_xml(blank_xml))

doc.save(target_path)
print("SUCCESS: 1-col kop, tight 1st title, enter before other titles applied!")
