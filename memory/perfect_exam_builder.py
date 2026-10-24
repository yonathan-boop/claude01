import docx
from docx.shared import Cm, Inches, Pt
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import zipfile

kop_path = r'C:\Users\yonat\Desktop\Soal Ujian Semester I T.P.2026-2027\SD SWASTA METHODIST.docx'
backup_path = r'C:\Users\yonat\.claude\projects\C--Users-yonat\memory\backups\Agama 1 Rio OK.raw_backup.docx'
target_path = r'C:\Users\yonat\Desktop\Soal Ujian Semester I T.P.2026-2027\Agama\Agama 1 Rio OK.docx'

# Extract logo image
with zipfile.ZipFile(kop_path, 'r') as z:
    img_data = z.read('word/media/image1.png')

doc_kop = docx.Document(kop_path)
tbl_xml = doc_kop.tables[0]._tbl.xml
# Map image embed cleanly to rIdLogo
tbl_xml = tbl_xml.replace('r:embed="rId5"', 'r:embed="rIdLogo"')

doc = docx.Document(backup_path)

def make_sectPr(cols=1):
    col_str = f'<w:cols {nsdecls("w")} w:num="{cols}" w:space="284"/>' if cols > 1 else f'<w:cols {nsdecls("w")} w:space="720"/>'
    return parse_xml(f'''
    <w:sectPr {nsdecls("w")}>
      <w:type w:val="continuous"/>
      <w:pgSz w:w="12240" w:h="18720"/>
      <w:pgMar w:top="283" w:right="567" w:bottom="340" w:left="567" w:header="720" w:footer="720" w:gutter="0"/>
      {col_str}
      <w:docGrid w:linePitch="360"/>
    </w:sectPr>
    ''')

def make_empty_p():
    return parse_xml(f'''
    <w:p {nsdecls("w")}>
      <w:pPr>
        <w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>
      </w:pPr>
    </w:p>
    ''')

def set_ind(p, kind='soal'):
    pPr = p._p.get_or_add_pPr()
    # Remove existing ind
    for el in pPr.xpath('.//w:ind'):
        pPr.remove(el)

    if kind == 'judul':
        # Remove any numPr from headings so it NEVER inherits list hanging indents!
        for el in pPr.xpath('.//w:numPr'):
            pPr.remove(el)
        pPr.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="0" w:firstLine="0" w:hanging="0"/>'))
    elif kind == 'soal':
        pPr.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="369" w:hanging="369"/>'))
    elif kind == 'opsi':
        pPr.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="737" w:hanging="369"/>'))
    elif kind == 'line':
        for el in pPr.xpath('.//w:numPr'):
            pPr.remove(el)
        pPr.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="369" w:firstLine="0" w:hanging="0"/>'))

# 1. Remove raw header lines
p0 = doc.paragraphs[0]._p
p1 = doc.paragraphs[1]._p
p0.getparent().remove(p0)
p1.getparent().remove(p1)

# Clean text
for p in doc.paragraphs:
    t = p.text
    if '\t' in t or '  ' in t:
        cleaned = t.replace('\t', ' ').strip()
        while '  ' in cleaned:
            cleaned = cleaned.replace('  ', ' ')
        if cleaned:
            p.text = cleaned

# 2. Insert Kop Table at top
new_tbl = parse_xml(tbl_xml)
doc.paragraphs[0]._p.addprevious(new_tbl)

# Fill Kop Table values and CLEAN all cell paragraph indents
t = doc.tables[0]
for r in t.rows:
    for c in r.cells:
        for p in c.paragraphs:
            pPr = p._p.get_or_add_pPr()
            for el in pPr.xpath('.//w:ind'):
                pPr.remove(el)
            for el in pPr.xpath('.//w:numPr'):
                pPr.remove(el)
            # Ensure cell paragraphs have 0 indent
            pPr.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="0" w:right="0" w:firstLine="0" w:hanging="0"/>'))

for r_i, val in [(0, 'Agama'), (1, 'Selasa, 1 Desember 2026'), (2, 'IA SD')]:
    cell_p = t.cell(r_i, 4).paragraphs[0]
    cell_p.text = ''
    run = cell_p.add_run(val)
    run.font.name = 'Tahoma'
    run.font.size = Pt(7)
    run.font.bold = True

# 3. Section 0: Kop Surat (1 col, continuous)
sep_sectPr = make_sectPr(cols=1)
sep_p_xml = f'''
<w:p {nsdecls("w")}>
  <w:pPr>
    <w:spacing w:before="0" w:after="0" w:line="20" w:lineRule="exact"/>
    <w:rPr>
      <w:sz w:val="2"/>
      <w:szCs w:val="2"/>
    </w:rPr>
  </w:pPr>
</w:p>
'''
sep_p = parse_xml(sep_p_xml)
sep_p.find(f'{{{nsdecls("w").split("=")[1].strip(chr(34))}}}pPr').append(sep_sectPr)
doc.paragraphs[0]._p.addprevious(sep_p)

# doc.paragraphs[0] is sep_p (1-pt)
# doc.paragraphs[1] is 'Pilihan Ganda CP 1 (100%)'
first_title = doc.paragraphs[1]
first_title.paragraph_format.space_before = Pt(0)
first_title.paragraph_format.space_after = Pt(2)
set_ind(first_title, 'judul')

# Ensure enter before 'CP 2 (100%)' and 'CP 3 (100%)'
for idx in range(len(doc.paragraphs)):
    p = doc.paragraphs[idx]
    txt = p.text.strip()
    if txt in ['CP 2 (100%)', 'CP 3 (100%)']:
        set_ind(p, 'judul')
        prev_p = doc.paragraphs[idx - 1]
        if prev_p.text.strip() != '':
            p._p.addprevious(make_empty_p())

# Format all PG content (120 items = 30 questions * (1 soal + 3 opsi))
pg_counter = 0
in_pg = False
for p in doc.paragraphs:
    txt = p.text.strip()
    if 'Pilihan Ganda CP 1' in txt:
        in_pg = True
        continue
    if 'Isian' in txt:
        in_pg = False
        break
    if in_pg and txt and not txt.startswith('CP '):
        if pg_counter % 4 == 0:
            set_ind(p, 'soal')
        else:
            set_ind(p, 'opsi')
        pg_counter += 1

print(f"Total PG items formatted: {pg_counter}")

# Find Isian heading
isian_p = None
for p in doc.paragraphs:
    if 'Isian CP 4' in p.text:
        isian_p = p
        break

# The last option of PG ends Section 1 (2 cols PG)
curr = isian_p._p.getprevious()
while curr is not None and curr.text is not None and curr.text.strip() == '':
    curr = curr.getprevious()
curr.get_or_add_pPr().append(make_sectPr(cols=2))

# Section 2: Enter 1 col + Isian heading 1 col
enter_1col = make_empty_p()
isian_p._p.addprevious(enter_1col)
set_ind(isian_p, 'judul')
isian_p.paragraph_format.space_before = Pt(0)
isian_p.paragraph_format.space_after = Pt(2)
isian_p._p.get_or_add_pPr().append(make_sectPr(cols=1))

# Find Esai heading
esai_p = None
for p in doc.paragraphs:
    if 'Esai (60%)' in p.text:
        esai_p = p
        break

# Format isian questions in 2 cols
curr_isian = isian_p._p.getnext()
while curr_isian is not None and curr_isian != esai_p._p:
    t_nodes = curr_isian.xpath('.//w:t')
    txt = ''.join(n.text for n in t_nodes if n.text).strip()
    if txt and 'Esai' not in txt:
        pPr = curr_isian.get_or_add_pPr()
        for el in pPr.xpath('.//w:ind'):
            pPr.remove(el)
        pPr.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="369" w:hanging="369"/>'))
    curr_isian = curr_isian.getnext()

# Last isian ends Section 3 (2 cols Isian)
last_isian = esai_p._p.getprevious()
while last_isian is not None and last_isian.text is not None and last_isian.text.strip() == '':
    last_isian = last_isian.getprevious()
last_isian.get_or_add_pPr().append(make_sectPr(cols=2))

# Section 4: Esai (1 col)
enter_esai = make_empty_p()
esai_p._p.addprevious(enter_esai)
set_ind(esai_p, 'judul')
esai_p.paragraph_format.space_before = Pt(0)
esai_p.paragraph_format.space_after = Pt(2)

# Format Esai questions and answer lines
curr_esai = esai_p._p.getnext()
while curr_esai is not None:
    t_nodes = curr_esai.xpath('.//w:t')
    txt = ''.join(n.text for n in t_nodes if n.text).strip()
    if txt:
        pPr = curr_esai.get_or_add_pPr()
        for el in pPr.xpath('.//w:ind'):
            pPr.remove(el)
        if txt.startswith('Jb:') or '____' in txt:
            for el in pPr.xpath('.//w:numPr'):
                pPr.remove(el)
            pPr.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="369" w:firstLine="0" w:hanging="0"/>'))
        else:
            pPr.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="369" w:hanging="369"/>'))
    curr_esai = curr_esai.getnext()

# Document-level sectPr (1 col, continuous)
body_sectPr = doc.sections[-1]._sectPr
final_sectPr = make_sectPr(cols=1)
parent = body_sectPr.getparent()
parent.replace(body_sectPr, final_sectPr)

doc.save(target_path)

# Inject logo image & relationship
with zipfile.ZipFile(target_path, 'r') as z_in:
    file_map = {name: z_in.read(name) for name in z_in.namelist()}

file_map['word/media/image1.png'] = img_data

rels_xml = file_map['word/_rels/document.xml.rels'].decode('utf-8')
if 'rIdLogo' not in rels_xml:
    rel_tag = '<Relationship Id="rIdLogo" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/image1.png"/>'
    rels_xml = rels_xml.replace('</Relationships>', rel_tag + '</Relationships>')
    file_map['word/_rels/document.xml.rels'] = rels_xml.encode('utf-8')

ct_xml = file_map['[Content_Types].xml'].decode('utf-8')
if 'Extension="png"' not in ct_xml:
    png_tag = '<Default Extension="png" ContentType="image/png"/>'
    ct_xml = ct_xml.replace('</Types>', png_tag + '</Types>')
    file_map['[Content_Types].xml'] = ct_xml.encode('utf-8')

with zipfile.ZipFile(target_path, 'w', zipfile.ZIP_DEFLATED) as z_out:
    for name, data in file_map.items():
        z_out.writestr(name, data)

print("SUCCESS: Full build completed with complete clean indents!")
